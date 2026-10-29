"""Normalize Lingxing read-only exports and the local shipment ledger for the dashboard."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import math
import os
import statistics
from collections import defaultdict
from pathlib import Path

DEFAULT_SCENARIO_PATH = Path(
    os.environ.get(
        "PMC_SCENARIO_PATH",
        str(Path(__file__).resolve().parents[1] / "state" / "scenarios.json"),
    )
)
from typing import Any


def config_root() -> Path:
    explicit = os.getenv("AMAZON_OPS_CONFIG_DIR")
    if explicit:
        return Path(explicit).expanduser().resolve()
    base = os.getenv("AMAZON_OPS_DATA_DIR")
    if not base:
        base = str(Path(os.getenv("LOCALAPPDATA") or Path.home() / ".local" / "share") / "amazon-ops-cockpit")
    return Path(base).expanduser().resolve() / "config"


def config_file(name: str) -> Path:
    external = config_root() / name
    return external if external.exists() else Path(__file__).parents[1] / "data" / name


def number(value: Any, default: float = 0.0) -> float:
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return default
    return parsed if math.isfinite(parsed) else default


def integer(value: Any, default: int = 0) -> int:
    return int(round(number(value, default)))


def unique(values: list[Any]) -> list[str]:
    return list(dict.fromkeys(str(value).strip() for value in values if str(value or "").strip()))


def clean_text(value: Any, fallback: str = "") -> str:
    text = str(value or "").strip()
    if not text or "�" in text:
        return fallback
    return text


def repeated(value: float, count: int) -> list[float]:
    return [max(0.0, float(value)) for _ in range(count)]


def history_from_sales_totals(sales: dict[str, Any]) -> list[float]:
    """Expand cumulative 14/30/60-day totals into auditable 20-day segments."""
    total14 = max(0.0, number(sales.get("salesTotal14")))
    total30 = max(total14, number(sales.get("salesTotal30")))
    total60 = max(total30, number(sales.get("salesTotal60")))
    avg_1_14 = total14 / 14
    avg_15_30 = max(0.0, total30 - total14) / 16
    avg_31_60 = max(0.0, total60 - total30) / 30
    # Chronological order: days 60..41, 40..21, 20..1.
    far = repeated(avg_31_60, 20)
    middle = repeated(avg_31_60, 10) + repeated(avg_15_30, 10)
    recent = repeated(avg_15_30, 6) + repeated(avg_1_14, 14)
    return far + middle + recent


def history_from_listing_totals(seven: float, fourteen: float, thirty: float) -> list[float]:
    total7 = max(0.0, seven)
    total14 = max(total7, fourteen)
    total30 = max(total14, thirty)
    # No 60-day value at Listing level: extend days 15-30 backward as an explicit fallback.
    older_daily = max(0.0, total30 - total14) / 16
    total60 = total30 + older_daily * 30
    return history_from_sales_totals(
        {"salesTotal14": total14, "salesTotal30": total30, "salesTotal60": total60}
    )


def fba_aged_buckets(info: dict[str, Any] | None) -> dict[str, int] | None:
    if not isinstance(info, dict):
        return None
    d91 = integer(info.get("invAge91To180Days"))
    d181 = integer(info.get("invAge181To270Days"))
    d271 = integer(info.get("invAge271To365Days"))
    d365 = integer(info.get("invAge365PlusDays"))
    return {
        "over90": d91 + d181 + d271 + d365,
        "over180": d181 + d271 + d365,
        "over270": d271 + d365,
        "over365": d365,
    }


def normalize_query_record(record: dict[str, Any]) -> dict[str, Any]:
    display = record.get("displayInfo") or {}
    basic = record.get("basicInfo") or {}
    amazon = record.get("amazonQuantityInfo") or {}
    scm = record.get("scmQuantityInfo") or {}
    sales = record.get("salesInfo") or {}
    suggest = record.get("suggestInfo") or {}
    relations = basic.get("relationListing") or []
    relation = relations[0] if relations else {}
    products = display.get("productList") or []
    product = products[0] if products else {}
    asin_list = display.get("asinList") or []
    asin = clean_text(display.get("asin")) or clean_text(relation.get("asin")) or (asin_list[0] if asin_list else "")
    local = integer(scm.get("scQuantityLocalValid"))
    overseas = integer(scm.get("scQuantityOverseaValid"))
    procurement = integer(scm.get("scQuantityPurchasePlan")) + integer(scm.get("scQuantityPurchaseShipping"))
    return {
        "asin": asin,
        "storeId": str(relation.get("storeId") or ""),
        "sku": clean_text(product.get("sku")),
        "productName": clean_text(product.get("productName")),
        "msku": clean_text(relation.get("msku")),
        "fnsku": clean_text(relation.get("fnsku")),
        "purchaseLeadDays": integer(display.get("daysPurchase"), 0),
        "fbaAvailable": integer(amazon.get("afnFulfillableQuantity")),
        "fbaReserved": integer(amazon.get("afnReservedQuantity")),
        "fbaInbound": integer(amazon.get("amazonQuantityShipping")),
        "localAvailable": local,
        "overseasAvailable": overseas,
        "sourceAvailable": local + overseas,
        "procurementOpen": procurement,
        "procurementPlan": integer(scm.get("scQuantityPurchasePlan")),
        "procurementShipping": integer(scm.get("scQuantityPurchaseShipping")),
        "agedBuckets": fba_aged_buckets(display.get("fbaAgedInfo")),
        "history": history_from_sales_totals(sales),
        "salesInfo": sales,
        "syncTime": clean_text(basic.get("syncTime")),
        "imageUrl": clean_text(display.get("smallImageUrl")),
        "itemName": clean_text(display.get("itemName")),
        "relations": relations,
        "products": products,
        "lingxingSuggestion": {
            "estimatedDaily": number(suggest.get("estimatedSaleAvgQuantity")),
            "availableSaleDays": integer(suggest.get("availableSaleDays")),
            "fbaAvailableSaleDays": integer(suggest.get("availableSaleDaysFba")),
            "fbaAndShippingDays": integer(suggest.get("availableSaleDaysFbaAndShipping")),
            "stockoutDate": clean_text(suggest.get("outStockDate")),
            "purchaseStockoutDate": clean_text(suggest.get("outStockDatePurchase")),
            "suggestSendDate": clean_text(suggest.get("sugDateSendLocal")),
            "suggestSendQty": integer(suggest.get("quantitySugLocalToFba")),
            "suggestPurchaseDate": clean_text(suggest.get("sugDatePurchase")),
            "suggestPurchaseQty": integer(suggest.get("quantitySugPurchase")),
            "restockStatus": integer((record.get("extInfo") or {}).get("restockStatus")),
        },
    }


def normalize_product_performance_row(record: dict[str, Any]) -> dict[str, Any]:
    """Normalize Lingxing product-performance output without merging inventory layers."""
    sids = record.get("sids") or []
    asin_items = record.get("asins") or []
    asin = clean_text(record.get("asin"))
    if not asin and asin_items:
        asin = clean_text((asin_items[0] or {}).get("asin"))
    sid = str((sids[0] if sids else "") or "")
    if not sid and asin_items:
        sid = str((asin_items[0] or {}).get("sid") or "")
    return {
        "asin": asin,
        "storeId": sid,
        "date": clean_text(record.get("rdate"))[:10],
        "sales": number(record.get("volume")),
        "inventoryLayers": {
            "sellable": integer(record.get("afn_fulfillable_quantity")),
            "reservedCustomerOrders": integer(record.get("reserved_customerorders")),
            "fcTransfers": integer(record.get("reserved_fc_transfers")),
            "fcProcessing": integer(record.get("reserved_fc_processing")),
            "receiving": integer(record.get("afn_inbound_receiving_quantity")),
            "shippedRaw": integer(record.get("afn_inbound_shipped_quantity")),
        },
    }


def product_performance_history(rows: list[dict[str, Any]], as_of: str) -> list[dict[str, Any]]:
    """Build 60 chronological calendar days while preserving missing stock evidence as None."""
    end = dt.date.fromisoformat(as_of)
    start = end - dt.timedelta(days=59)
    by_date = {clean_text(row.get("date"))[:10]: row for row in rows if clean_text(row.get("date"))}
    history: list[dict[str, Any]] = []
    for offset in range(60):
        date_text = (start + dt.timedelta(days=offset)).isoformat()
        row = by_date.get(date_text)
        layers = (row or {}).get("inventoryLayers") or {}
        history.append(
            {
                "date": date_text,
                "sales": number(row["sales"]) if row and row.get("sales") is not None else None,
                "fbaSellable": integer(layers["sellable"]) if layers.get("sellable") is not None else None,
                "inboundReceiving": integer(layers["receiving"]) if layers.get("receiving") is not None else None,
            }
        )
    return history


def serialize_history(history: list[Any]) -> list[Any]:
    """Round numeric demand values while preserving enriched day evidence."""
    serialized: list[Any] = []
    for item in history:
        if isinstance(item, dict):
            serialized.append(
                {
                    **item,
                    "sales": round(number(item["sales"]), 4) if item.get("sales") is not None else None,
                    "fbaSellable": (
                        None if item.get("fbaSellable") is None else integer(item.get("fbaSellable"))
                    ),
                    "inboundReceiving": (
                        None
                        if item.get("inboundReceiving") is None
                        else integer(item.get("inboundReceiving"))
                    ),
                }
            )
        else:
            serialized.append(round(number(item), 4) if item is not None else None)
    return serialized


def load_products(path: Path | None) -> dict[str, dict[str, str]]:
    if not path or not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    result: dict[str, dict[str, str]] = {}
    for group in data.get("groups", []):
        for product in group.get("products", []):
            result[str(product.get("asin") or "")] = {
                "name": str(product.get("name") or ""),
                "group": str(group.get("name") or ""),
            }
    return result


def load_transit_policy(path: Path) -> dict[str, Any]:
    fallback = {
        "normalMode": "海运", "normalDays": 60,
        "urgentMode": "卡航", "urgentDays": 30,
        "modes": {"卡航": 30, "海运": 60, "铁路": 40, "空运": 10},
    }
    if not path.exists():
        return fallback
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return fallback
    return {**fallback, **data}


def resolve_transit(
    asin: str,
    sid: str,
    skus: list[str],
    policy: dict[str, Any],
    medians: dict[str, int],
) -> dict[str, Any]:
    """Normal lane for routine replenishment, urgent lane when a batch runs late."""
    modes = policy.get("modes") or {}
    normal_mode = str(policy.get("normalMode") or "海运")
    normal_days = integer(policy.get("normalDays"), 0) or integer(modes.get(normal_mode), 60) or 60
    urgent_mode = str(policy.get("urgentMode") or "卡航")
    urgent_days = integer(policy.get("urgentDays"), 0) or integer(modes.get(urgent_mode), 30) or 30
    evidence = medians.get(asin)
    evidence_note = f"；历史货件中位 {evidence} 天仅作参考" if evidence else ""
    source = f"{normal_mode} {normal_days} 天 / {urgent_mode} {urgent_days} 天{evidence_note}"

    for sku in skus or []:
        entry = (policy.get("bySku") or {}).get(sku)
        if isinstance(entry, dict):
            normal_mode = str(entry.get("normalMode") or normal_mode)
            normal_days = integer(entry.get("normalDays"), 0) or integer(modes.get(normal_mode), normal_days)
            urgent_mode = str(entry.get("urgentMode") or urgent_mode)
            urgent_days = integer(entry.get("urgentDays"), 0) or integer(modes.get(urgent_mode), urgent_days)
            source = f"SKU 指定 {short_transit(normal_mode, normal_days)} / {short_transit(urgent_mode, urgent_days)}{evidence_note}"
            break

    if not source.startswith("SKU"):
        entry = (policy.get("byStore") or {}).get(sid)
        if isinstance(entry, dict):
            normal_mode = str(entry.get("normalMode") or normal_mode)
            normal_days = integer(entry.get("normalDays"), 0) or integer(modes.get(normal_mode), normal_days)
            urgent_mode = str(entry.get("urgentMode") or urgent_mode)
            urgent_days = integer(entry.get("urgentDays"), 0) or integer(modes.get(urgent_mode), urgent_days)
            source = f"店铺策略 {short_transit(normal_mode, normal_days)} / {short_transit(urgent_mode, urgent_days)}{evidence_note}"

    return {
        "normalMode": normal_mode,
        "normalDays": normal_days,
        "urgentMode": urgent_mode,
        "urgentDays": urgent_days,
        "source": source,
    }


def short_transit(mode: str, days: int) -> str:
    return f"{mode} {days} 天"


def _date_text(value: Any) -> str:
    if isinstance(value, (dt.date, dt.datetime)):
        return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()
    return clean_text(value)[:10]


def shipments_from_rows(
    rows: list[dict[str, Any]], as_of: str, source_name: str
) -> tuple[dict[tuple[str, str], list[dict[str, Any]]], dict[str, int], dict[str, Any]]:
    inbound: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    transit_values: dict[str, list[int]] = defaultdict(list)
    as_of_date = dt.date.fromisoformat(as_of)
    active_statuses = {"SHIPPED", "IN_TRANSIT", "CHECKED_IN"}
    planned_statuses = {"WORKING", "READY_TO_SHIP"}
    # FBA exports merge shipment-level cells across the multiple ASIN rows of one
    # shipment. Carry the parent row's fields into those continuation rows; a row
    # with a new shipment id starts a new carried block.
    shipment_fields = (
        "货件单号", "店铺", "货件状态", "货件实际状态", "到货时间", "实际发货时间",
        "SHIPPED状态变更时间", "送达时段(开始时间)", "送达时段(结束时间)", "运输方式", "更新时间",
    )
    carried: dict[str, Any] = {}
    active_rows = 0
    planned_rows = 0
    receiving_rows = 0
    unknown_eta_rows = 0
    unverified_rows = 0
    for raw_row in rows:
        row = dict(raw_row)
        marker = clean_text(row.get("货件单号") or row.get("shipment_id"))
        if marker and marker != clean_text(carried.get("货件单号")):
            carried = {}
        for field in shipment_fields:
            value = row.get(field)
            if value not in (None, ""):
                carried[field] = value
            elif field in carried:
                row[field] = carried[field]
        asin = clean_text(row.get("ASIN") or row.get("asin"))
        if not asin:
            continue
        store = clean_text(row.get("店铺") or row.get("store"))
        status = clean_text(row.get("货件状态") or row.get("actual_status") or row.get("shipment_status")).upper()
        if status in planned_statuses:
            planned_rows += 1
        if status == "RECEIVING":
            receiving_rows += 1
        shipped_date = _date_text(row.get("实际发货时间") or row.get("SHIPPED状态变更时间"))
        arrival_date = _date_text(row.get("到货时间"))
        historical_days = 0
        try:
            if shipped_date and arrival_date:
                historical_days = (dt.date.fromisoformat(arrival_date) - dt.date.fromisoformat(shipped_date)).days
        except ValueError:
            historical_days = 0
        historical_days = historical_days or integer(row.get("shipped_to_eta_start_days"))
        if 5 <= historical_days <= 90:
            transit_values[asin].append(historical_days)
        if status not in active_statuses:
            continue
        declared = integer(row.get("申报量") or row.get("declared_qty"))
        shipped = integer(row.get("已发货") or row.get("shipped_qty"))
        received = integer(row.get("签收量") or row.get("received_qty"))
        # A shipment that is still IN_TRANSIT with zero shipped units is contradictory:
        # keep the declared quantity visible but never promote it to confirmed supply.
        quantity_verified = shipped > 0
        qty = max(0, (shipped if quantity_verified else declared) - received)
        if qty <= 0:
            continue
        active_rows += 1
        if not quantity_verified:
            unverified_rows += 1
        eta_start = _date_text(row.get("送达时段(开始时间)") or row.get("eta_start"))
        eta_end = _date_text(row.get("送达时段(结束时间)") or row.get("eta_end"))
        valid_start = ""
        valid_end = ""
        for label, candidate in (("start", eta_start), ("end", eta_end)):
            try:
                parsed = dt.date.fromisoformat(candidate)
            except ValueError:
                continue
            if parsed >= as_of_date:
                if label == "start":
                    valid_start = parsed.isoformat()
                else:
                    valid_end = parsed.isoformat()
        if valid_start and not valid_end:
            valid_end = valid_start
        if valid_end and not valid_start:
            valid_start = valid_end
        if not valid_end:
            unknown_eta_rows += 1
        shipment_id = clean_text(row.get("货件单号") or row.get("shipment_id"))
        inbound[(store, asin)].append(
            {
                "date": valid_end or None,
                "etaStart": valid_start or None,
                "etaEnd": valid_end or None,
                "qty": qty,
                "confirmed": bool(valid_end) and quantity_verified and bool(shipment_id),
                "quantityVerified": quantity_verified,
                "shipment": shipment_id or "货件记录（缺货件号）",
                "status": status,
                "etaWindow": " ~ ".join(value for value in (eta_start, eta_end) if value),
                "source": f"FBA货件台账 {source_name}",
                "sku": clean_text(row.get("SKU") or row.get("sku")),
                "msku": clean_text(row.get("MSKU") or row.get("msku")),
                "quantitySource": "已发货-签收量" if quantity_verified else "申报量-签收量（已发货为0，仅作参考）",
            }
        )
    medians = {asin: int(round(statistics.median(values))) for asin, values in transit_values.items() if values}
    last_update = ""
    for raw_row in rows:
        candidate = _date_text(raw_row.get("更新时间") or raw_row.get("updated_at"))
        if candidate and candidate > last_update:
            last_update = candidate
    reminders: list[str] = []
    if unknown_eta_rows:
        reminders.append(f"{unknown_eta_rows}条活动货件缺少未来ETA，请更新FBA货件文件")
    if unverified_rows:
        reminders.append(f"{unverified_rows}条活动货件已发货为0，已按参考数量处理且不计入确认在途")
    metadata = {
        "status": "loaded",
        "sourceName": source_name,
        "rows": len(rows),
        "activeRows": active_rows,
        "plannedRows": planned_rows,
        "receivingRows": receiving_rows,
        "unknownEtaRows": unknown_eta_rows,
        "unverifiedRows": unverified_rows,
        "lastUpdatedAt": last_update,
        "reminder": "；".join(reminders),
    }
    return inbound, medians, metadata


def load_shipments(
    path: Path | None, as_of: str
) -> tuple[dict[tuple[str, str], list[dict[str, Any]]], dict[str, int], dict[str, Any]]:
    inbound: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    if not path or not path.exists():
        return inbound, {}, {
            "status": "missing",
            "sourceName": path.name if path else "FBA货件.xlsx",
            "rows": 0,
            "activeRows": 0,
            "plannedRows": 0,
            "receivingRows": 0,
            "unknownEtaRows": 0,
            "unverifiedRows": 0,
            "lastUpdatedAt": "",
            "reminder": "未找到ETA文件，请补充FBA货件.xlsx后重新生成看板",
        }
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        required = {"货件单号", "店铺", "货件状态", "ASIN", "已发货", "签收量"}
        candidate = []
        for current in workbook:
            header = [clean_text(value) for value in next(current.iter_rows(values_only=True), [])]
            if required.issubset(set(header)) and ("送达时段(结束时间)" in header or "eta_end" in header):
                candidate.append(current)
        if len(candidate) != 1:
            workbook.close()
            raise ValueError(f"ETA workbook requires exactly one complete shipment sheet; found {len(candidate)}")
        sheet = candidate[0]
        iterator = sheet.iter_rows(values_only=True)
        headers = [clean_text(value) for value in next(iterator)]
        rows = [dict(zip(headers, values)) for values in iterator]
        workbook.close()
    else:
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    return shipments_from_rows(rows, as_of, path.name)


def exact_daily_history(sales_row: dict[str, Any], as_of: str) -> list[float]:
    date_collect = sales_row.get("dateCollect") or {}
    anchor_text = clean_text((sales_row.get("rangeAvg") or {}).get("anchorDate"), as_of)
    try:
        anchor = min(dt.date.fromisoformat(anchor_text), dt.date.fromisoformat(as_of))
    except ValueError:
        anchor = dt.date.fromisoformat(as_of)
    start = anchor - dt.timedelta(days=59)
    return [number(date_collect[key]) if date_collect.get(key) is not None else None
            for offset in range(60) for key in [(start + dt.timedelta(days=offset)).isoformat()]]


def load_scenarios(path: Path | None) -> dict[str, Any]:
    if not path or not path.exists():
        return {"version": 1, "updatedAt": None, "scenarios": []}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise ValueError(f"scenario file is invalid: {path}") from error
    if not isinstance(data, dict) or not isinstance(data.get("scenarios"), list):
        raise ValueError(f"scenario file has no scenarios list: {path}")
    for item in data["scenarios"]:
        if not isinstance(item, dict) or not item.get("store") or not item.get("asin"):
            raise ValueError(f"scenario entry missing store/asin: {path}")
    return data


def attach_scenarios(records: list[dict[str, Any]], scenario_data: dict[str, Any]) -> int:
    index: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for scenario in scenario_data.get("scenarios", []):
        key = (str(scenario.get("store") or ""), str(scenario.get("asin") or ""))
        if not key[1]:
            continue
        index.setdefault(key, []).append(scenario)
    attached = 0
    for record in records:
        items = index.get((record["store"], record["asin"]), [])
        record["scenarios"] = items
        default = next((item for item in items if item.get("isDefault")), None)
        record["activeScenarioId"] = (default or (items[0] if items else {})).get("id") if items else None
        attached += len(items)
    return attached


def allocate_shared_sku_stock(records: list[dict[str, Any]]) -> int:
    """Split a shared SKU's local/overseas stock by each store's 30-day unit share."""
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        if record.get("listingStatus") != 1:
            continue
        skus = record.get("skus") or ([record["sku"]] if record.get("sku") else [])
        for sku in set(skus):
            groups[sku].append(record)
    changed = 0
    for sku, group in groups.items():
        if len(group) < 2:
            continue
        units = [max(0.0, number(record.get("units30"))) for record in group]
        total_units = sum(units)
        shares = [unit / total_units for unit in units] if total_units > 0 else [1 / len(group)] * len(group)
        for record, share in zip(group, shares):
            layers = record["inventoryLayers"]
            totals = {
                "local": layers.get("local", 0),
                "overseas": layers.get("overseas", 0),
                "procurementOpen": layers.get("procurementOpen", 0),
            }
            record["stockAllocation"] = {
                "sku": sku,
                "share": round(share, 4),
                "stores": len(group),
                "skuTotals": totals,
            }
            for key in totals:
                layers[key] = round(totals[key] * share, 2)
            record["sourceAvailable"] = round(number(record.get("sourceAvailable")) * share, 2)
            changed += 1
    return changed


def build_snapshot(
    export_path: Path,
    shipment_path: Path | None = None,
    product_path: Path | None = None,
    scenario_path: Path | None = None,
) -> dict[str, Any]:
    export = json.loads(export_path.read_text(encoding="utf-8"))
    as_of = str(export.get("as_of") or "")
    if not as_of:
        raise ValueError("Lingxing export requires as_of")
    dt.date.fromisoformat(as_of)
    tools = export["tools"]
    listings = tools["erp_listing"]["response"]["data"]["data"]["list"]
    stocks = tools["get_fba_stock_list"]["response"]["data"].get("list", [])
    query_rows = tools["query_fba_valid_list"]["response"]["data"]["data"].get("list", [])
    sales_rows = tools["platform_v2_page_list"]["response"]["data"]["data"].get("statisticsList", [])
    performance_tool = tools.get("query_product_performance_asin_lists") or {}
    performance_response = performance_tool.get("response") or {}
    performance_rows_raw = (
        ((performance_response.get("data") or {}).get("data") or {}).get("list") or []
    )
    products = load_products(product_path)
    carton_rule_path = config_file("carton-rules.json")
    store_names_path = config_file("store-names.json")
    store_names = json.loads(store_names_path.read_text(encoding="utf-8")) if store_names_path.exists() else {}
    carton_rules = json.loads(carton_rule_path.read_text(encoding="utf-8")) if carton_rule_path.exists() else {}
    transit_policy = load_transit_policy(config_file("transit-policy.json"))
    shipment_map, transit_medians, eta_file = load_shipments(shipment_path, as_of)

    listing_groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in listings:
        key = (str(row.get("store_id") or ""), str(row.get("asin") or ""))
        if key[1]:
            listing_groups[key].append(row)

    stock_groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in stocks:
        key = (str(row.get("sid") or row.get("group_by_sid") or ""), str(row.get("asin") or ""))
        if key[1]:
            stock_groups[key].append(row)

    query_by_key: dict[tuple[str, str], dict[str, Any]] = {}
    query_by_asin: dict[str, dict[str, Any]] = {}
    for row in query_rows:
        normalized = normalize_query_record(row)
        if normalized["asin"]:
            query_by_asin[normalized["asin"]] = normalized
        relations = normalized.get("relations") or [{}]
        for relation in relations:
            asin = clean_text(relation.get("asin")) or normalized["asin"]
            sid = str(relation.get("storeId") or normalized["storeId"] or "")
            if asin:
                query_by_key[(sid, asin)] = normalized

    sales_by_key: dict[tuple[str, str], dict[str, Any]] = {}
    for row in sales_rows:
        asins = row.get("platformProductId") or []
        stores = row.get("storeId") or []
        if asins:
            sales_by_key[(str(stores[0]) if stores else "", str(asins[0]))] = row

    performance_by_key: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in performance_rows_raw:
        normalized = normalize_product_performance_row(row)
        if normalized["asin"]:
            performance_by_key[(normalized["storeId"], normalized["asin"])].append(normalized)
    for rows in performance_by_key.values():
        rows.sort(key=lambda row: row.get("date") or "")

    records: list[dict[str, Any]] = []
    eta_covered = 0
    exact_sales = 0
    performance_sales = 0
    default_lead = 0
    eta_ledger_total = 0
    eta_live_total = 0
    eta_trimmed_total = 0
    eta_trimmed_asins = 0
    eta_orphan_asins = 0
    eta_orphan_total = 0
    for (sid, asin), group in listing_groups.items():
        query = query_by_key.get((sid, asin))
        stock_rows = stock_groups.get((sid, asin), [])
        product_config = products.get(asin, {})
        skus = unique([row.get("local_sku") for row in group] + ([query.get("sku")] if query else []))
        carton_rule = carton_rules.get(f"{sid}-{asin}", {})
        if carton_rule and carton_rule.get("sku") not in skus:
            carton_rule = {}
        mskus = unique([row.get("msku") for row in group] + ([query.get("msku")] if query else []))
        fnskus = unique([row.get("fnsku") for row in group] + ([query.get("fnsku")] if query else []))
        sales_row = sales_by_key.get((sid, asin)) or sales_by_key.get(("", asin))
        performance_rows = performance_by_key.get((sid, asin), [])
        if performance_rows:
            history = product_performance_history(performance_rows, as_of)
            forecast_source = "领星产品表现：店铺-ASIN最近60个完整日销量与库存约束"
            exact_sales += 1
            performance_sales += 1
        elif sales_row:
            history = exact_daily_history(sales_row, as_of)
            forecast_source = "领星销量统计：截至完整日的60天逐日销量"
            exact_sales += 1
        elif query:
            history = query["history"]
            forecast_source = "领星补货分析：14/30/60日累计拆分为近中远20日段"
        else:
            volume7 = sum(number(row.get("seven_volume")) for row in group)
            volume14 = sum(number(row.get("fourteen_volume")) for row in group)
            volume30 = sum(number(row.get("thirty_volume")) for row in group)
            history = history_from_listing_totals(volume7, volume14, volume30)
            forecast_source = "领星Listing：7/14/30日累计外推60日（低置信度）"

        listing_fba = sum(integer(row.get("afn_fulfillable_quantity")) for row in group)
        listing_reserved = sum(integer(row.get("afn_reserved_quantity")) for row in group)
        listing_working = sum(integer(row.get("afn_inbound_working_quantity")) for row in group)
        listing_shipped = sum(integer(row.get("afn_inbound_shipped_quantity")) for row in group)
        listing_receiving = sum(integer(row.get("afn_inbound_receiving_quantity")) for row in group)
        stock_fba = sum(integer(row.get("afn_fulfillable_quantity")) for row in stock_rows)
        stock_shipped = sum(integer(row.get("afn_inbound_shipped_quantity")) for row in stock_rows)
        stock_receiving = sum(integer(row.get("afn_inbound_receiving_quantity")) for row in stock_rows)
        stock_aged = fba_aged_buckets(
            {
                "invAge91To180Days": sum(integer(row.get("inv_age_91_to_180_days")) for row in stock_rows),
                "invAge181To270Days": sum(integer(row.get("inv_age_181_to_270_days")) for row in stock_rows),
                "invAge271To365Days": sum(integer(row.get("inv_age_271_to_365_days")) for row in stock_rows),
                "invAge365PlusDays": sum(integer(row.get("inv_age_365_plus_days")) for row in stock_rows),
            }
        ) if stock_rows else None
        latest_performance = performance_rows[-1] if performance_rows else None
        performance_layers = (latest_performance or {}).get("inventoryLayers") or {}
        # Historical product-performance stock is a day-level demand constraint,
        # not the authority for today's FBA quantity.
        fba_available = stock_fba if stock_rows else (query["fbaAvailable"] if query else listing_fba)
        fba_reserved = listing_reserved if group else (query["fbaReserved"] if query else 0)
        live_inbound = stock_shipped if stock_rows else (query["fbaInbound"] if query else listing_shipped)
        inventory_layers = {
            "sellable": fba_available,
            "reservedCustomerOrders": fba_reserved,
            "fcTransfers": performance_layers.get("fcTransfers", 0),
            "fcProcessing": performance_layers.get("fcProcessing", 0),
            "receiving": stock_receiving if stock_rows else listing_receiving,
            "shippedRaw": live_inbound,
            "local": query.get("localAvailable", 0) if query else 0,
            "overseas": query.get("overseasAvailable", 0) if query else 0,
            "procurementOpen": query.get("procurementOpen", 0) if query else 0,
        }
        shipment_store = store_names.get(sid, "")
        inbound = [
            dict(item)
            for item in (
                shipment_map.get((shipment_store, asin), [])
                or shipment_map.get(("", asin), [])
            )
        ]
        eta_qty = sum(item["qty"] for item in inbound if item.get("confirmed"))
        unknown_ledger_qty = sum(item["qty"] for item in inbound if not item.get("confirmed"))
        residual_unknown = max(0, live_inbound - eta_qty - unknown_ledger_qty)
        if residual_unknown:
            inbound.append(
                {
                    "date": None,
                    "qty": residual_unknown,
                    "confirmed": False,
                    "shipment": "领星实时在途（ETA待匹配）",
                    "status": "在途",
                    "etaWindow": "",
                    "source": "领星补货/FBA库存",
                }
            )
        if any(item.get("confirmed") for item in inbound):
            eta_covered += 1
        ledger_qty = sum(item["qty"] for item in inbound)
        confirmed_qty = sum(item["qty"] for item in inbound if item.get("confirmed"))
        over_live_qty = max(0, ledger_qty - live_inbound)
        orphan_ledger_qty = ledger_qty if live_inbound <= 0 else 0
        eta_ledger_total += ledger_qty
        eta_live_total += live_inbound
        if over_live_qty:
            eta_trimmed_asins += 1
            eta_trimmed_total += over_live_qty
        if orphan_ledger_qty:
            eta_orphan_asins += 1
            eta_orphan_total += orphan_ledger_qty
        eta_reconciliation = {
            "ledgerQty": ledger_qty,
            "liveQty": live_inbound,
            "confirmedQty": confirmed_qty,
            "overLiveQty": over_live_qty,
            "orphanLedgerQty": orphan_ledger_qty,
        }

        aged = query.get("agedBuckets") if query else stock_aged
        revenue30 = sum(number(row.get("thirty_amount")) for row in group)
        revenue7 = sum(number(row.get("seven_amount")) for row in group)
        units30 = sum(number(row.get("thirty_volume")) for row in group)
        price = max((number(row.get("price")) for row in group), default=0)
        purchase_lead = query.get("purchaseLeadDays", 0) if query else 0
        if purchase_lead <= 0:
            purchase_lead = 40
            lead_source = "默认40天（领星未维护）"
            default_lead += 1
        else:
            lead_source = "领星补货建议 daysPurchase"
        transit_info = resolve_transit(
            asin, sid, skus, transit_policy, transit_medians
        )
        title = product_config.get("name") or clean_text(group[0].get("local_name")) or clean_text(group[0].get("item_name")) or asin
        store = store_names.get(sid) or product_config.get("group") or f"店铺 {sid}"
        lingxing = query.get("lingxingSuggestion") if query else None
        source_available = query.get("sourceAvailable", 0) if query else 0
        procurement_open = query.get("procurementOpen", 0) if query else 0
        gaps = []
        if not performance_rows and not sales_row:
            gaps.append("无60天逐日销量，已用领星累计销量拆段")
        if not query:
            gaps.append("无领星ASIN补货分析记录，源头库存按0保守处理")
        if not any(item.get("confirmed") for item in inbound) and live_inbound > 0:
            gaps.append("实时在途有数量但缺可靠ETA，不用于推迟断货日")
        if purchase_lead == 40 and lead_source.startswith("默认"):
            gaps.append("领星未维护采购天数，采用40天默认值")
        completeness_flags = {
            "skuMapping": bool(skus or mskus),
            "dailySales60": bool(performance_rows or sales_row),
            "inventoryBreakdown": bool(latest_performance),
            "replenishmentRecord": bool(query),
            "purchaseLead": not lead_source.startswith("默认"),
            "transitEvidence": asin in transit_medians,
            "etaEvidence": any(item.get("confirmed") for item in inbound) or live_inbound == 0,
            "agedInventory": aged is not None,
        }
        completeness = round(100 * sum(completeness_flags.values()) / len(completeness_flags))
        history_sales = sum(
            number(item.get("sales")) if isinstance(item, dict) else number(item)
            for item in history
        )
        evidence = [
            {"stage": "映射", "source": "领星 erp_listing", "fact": f"ASIN {asin} → SKU {', '.join(skus) or '未配对'} → MSKU {', '.join(mskus) or '—'}"},
            {"stage": "需求", "source": forecast_source, "fact": f"60天样本销量 {round(history_sales)} 件"},
            {"stage": "FBA", "source": "领星产品表现 / FBA库存", "fact": f"可售 {fba_available}，客户订单预留 {inventory_layers['reservedCustomerOrders']}，调仓 {inventory_layers['fcTransfers']}，Receiving {inventory_layers['receiving']}，Shipped {live_inbound}"},
            {"stage": "源头", "source": "领星补货建议 SCM 库存", "fact": f"本地 {query.get('localAvailable', 0) if query else 0}，海外仓 {query.get('overseasAvailable', 0) if query else 0}，采购未结 {procurement_open}"},
            {"stage": "时效", "source": f"{lead_source}；{transit_info['source']}", "fact": f"采购 {purchase_lead} 天，海运 {transit_info['normalDays']} 天 / 卡航 {transit_info['urgentDays']} 天，安全库存策略 15 天"},
        ]
        if carton_rule:
            evidence.append({"stage": "箱规", "source": carton_rule["source"], "fact": f"SKU {carton_rule['sku']}，单箱 {carton_rule['cartonQty']} 件"})
        if lingxing:
            evidence.append(
                {"stage": "交叉核验", "source": "领星原生补货建议", "fact": f"领星日均 {lingxing['estimatedDaily']}，建议发货日 {lingxing['suggestSendDate'] or '—'}，建议采购日 {lingxing['suggestPurchaseDate'] or '—'}"}
            )

        records.append(
            {
                "id": f"{sid}-{asin}",
                "storeId": sid,
                "store": store,
                "title": title,
                "asin": asin,
                "sku": skus[0] if skus else "未配对",
                "skus": skus,
                "msku": mskus[0] if mskus else "—",
                "mskus": mskus,
                "fnskus": fnskus,
                "imageUrl": clean_text(group[0].get("small_image_url")) or (query.get("imageUrl") if query else ""),
                "imageMark": "".join(part[:1] for part in title.split()[:2]).upper() or "SKU",
                "listingStatus": integer(group[0].get("status")),
                "listingStatusText": clean_text(group[0].get("status_text"), "在售" if integer(group[0].get("status")) == 1 else "非在售"),
                "history": serialize_history(history),
                "forecastSource": forecast_source,
                "fbaAvailable": fba_available,
                "fbaReserved": fba_reserved,
                "fbaInboundLive": live_inbound,
                "inventoryLayers": inventory_layers,
                "localAvailable": query.get("localAvailable", 0) if query else 0,
                "overseasAvailable": query.get("overseasAvailable", 0) if query else 0,
                "sourceAvailable": source_available,
                "procurementOpen": procurement_open,
                "procurementPlan": query.get("procurementPlan", 0) if query else 0,
                "procurementShipping": query.get("procurementShipping", 0) if query else 0,
                "inbound": inbound,
                "etaReconciliation": eta_reconciliation,
                "agedBuckets": aged,
                "price": round(price, 2),
                "revenue30": round(revenue30, 2),
                "revenue7": round(revenue7, 2),
                "units30": round(units30, 0),
                "purchaseLeadDays": purchase_lead,
                "cartonQty": int(carton_rule.get("cartonQty", 0)),
                "cartonSource": carton_rule.get("source", ""),
                "transitDays": transit_info["normalDays"],
                "transitMode": transit_info["normalMode"],
                "expressTransitDays": transit_info["urgentDays"],
                "expressMode": transit_info["urgentMode"],
                "transitEvidenceDays": transit_medians.get(asin, None),
                "safetyDays": 15,
                "cadenceDays": 7,
                "purchaseCadenceDays": 28,
                "leadTimeSource": lead_source,
                "transitSource": transit_info["source"],
                "lingxingSuggestion": lingxing,
                "dataCompleteness": completeness,
                "completenessFlags": completeness_flags,
                "dataGaps": gaps,
                "evidence": evidence,
                "sourceUpdatedAt": query.get("syncTime") if query else clean_text(group[0].get("last_referral_fee_preview_report_time")),
                "notes": gaps or ["关键链路字段完整，可按周滚动复核"],
            }
        )

    records.sort(key=lambda item: (item["listingStatus"] != 1, item["store"], item["asin"]))
    shared_allocated = allocate_shared_sku_stock(records)
    active_records = [record for record in records if record["listingStatus"] == 1]
    ranked = sorted(active_records, key=lambda record: record["revenue30"], reverse=True)
    revenue_total = sum(record["revenue30"] for record in ranked) or 1.0
    for index, record in enumerate(ranked):
        record["revenueRank"] = index + 1
        record["revenueShare"] = round(record["revenue30"] / revenue_total, 4)
        record["priorityTier"] = "P0" if index < 10 else ("P1" if index < 20 else "P2")
    top10_share = round(sum(record["revenue30"] for record in ranked[:10]) / revenue_total, 4)
    scenario_data = load_scenarios(scenario_path)
    scenario_count = attach_scenarios(records, scenario_data)
    warning_data = tools["settings_message_inventory_list"]["response"]["data"].get("data", {})
    purchase_response = tools["analytics_show_product_list"]["response"]["data"]
    metadata = {
        "asOf": as_of,
        "sourceAudit": export.get("source_audit") or {},
        "generatedAt": dt.datetime.now().strftime("%Y-%m-%d %H:%M CST"),
        "sourceMode": "领星 MCP 真实店铺快照",
        "sourceNote": f"已读取{len(listings)}条Listing并归并为{len(records)}个店铺-ASIN；补货分析{len(query_rows)}条、60天逐日销量{exact_sales}条。",
        "mcpStatus": "只读连接正常",
        "etaStatus": f"{eta_covered}/{len(records)} 个ASIN有可用于时间轴的确认ETA；其余在途不凭数量推迟断货日",
        "etaFile": eta_file,
        "transitPolicy": transit_policy,
        "revenue": {
            "total30": round(revenue_total, 2),
            "top10Share": top10_share,
            "activeRows": len(ranked),
        },
        "etaReconciliation": {
            "ledgerQty": eta_ledger_total,
            "liveQty": eta_live_total,
            "overLiveQty": eta_trimmed_total,
            "overLiveAsins": eta_trimmed_asins,
            "orphanLedgerQty": eta_orphan_total,
            "orphanAsins": eta_orphan_asins,
        },
        "writebackStatus": "仅草稿，不执行",
        "scenarios": {
            "path": str(scenario_path) if scenario_path else None,
            "updatedAt": scenario_data.get("updatedAt"),
            "count": scenario_count,
            "asins": len({(item.get("store"), item.get("asin")) for item in scenario_data.get("scenarios", [])}),
        },
        "stockAllocation": {
            "sharedSkuRecords": shared_allocated,
            "basis": "按各店铺近30天销量占比分摊共享SKU的本地/海外/采购未结",
        },
        "scope": {
            "listingRows": len(listings),
            "storeAsins": len(records),
            "uniqueAsins": len({record["asin"] for record in records}),
            "restockRows": len(query_rows),
            "dailySalesRows": exact_sales,
            "productPerformanceRows": len(performance_rows_raw),
            "productPerformanceAsins": performance_sales,
            "warningRows": integer(warning_data.get("total")),
            "etaCoveredRows": eta_covered,
            "defaultPurchaseLeadRows": default_lead,
        },
        "capabilities": [
            {"toolId": "erp_listing", "label": "ASIN / MSKU / SKU 映射", "state": "live", "rows": len(listings)},
            {"toolId": "get_fba_stock_list", "label": "FBA库存与库龄", "state": "live", "rows": len(stocks)},
            {"toolId": "query_fba_valid_list", "label": "补货、采购周期与SCM库存", "state": "live", "rows": len(query_rows)},
            {"toolId": "platform_v2_page_list", "label": "60天逐日销量", "state": "live", "rows": exact_sales},
            {"toolId": "query_product_performance_asin_lists", "label": "逐日销量与FBA库存分层", "state": "live" if performance_rows_raw else "missing", "rows": len(performance_rows_raw)},
            {"toolId": "settings_message_inventory_list", "label": "库存预警交叉核验", "state": "live", "rows": integer(warning_data.get("total"))},
            {"toolId": "analytics_show_product_list", "label": "采购报表明细", "state": "permission-missing" if purchase_response.get("code") == 0 else "live", "rows": 0},
            {"toolId": "local_fba_shipments", "label": "FBA货件ETA台账", "state": "local", "rows": eta_covered},
        ],
    }
    return {"metadata": metadata, "records": records}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--export", type=Path, required=True)
    parser.add_argument("--shipments", type=Path)
    parser.add_argument("--products", type=Path)
    parser.add_argument("--output-js", type=Path, required=True)
    parser.add_argument("--output-json", type=Path)
    parser.add_argument("--scenarios", type=Path, default=DEFAULT_SCENARIO_PATH)
    args = parser.parse_args()
    snapshot = build_snapshot(args.export, args.shipments, args.products, args.scenarios)
    args.output_js.parent.mkdir(parents=True, exist_ok=True)
    serialized = json.dumps(snapshot, ensure_ascii=False, indent=2)
    args.output_js.write_text(f"export const inventorySnapshot = {serialized};\n", encoding="utf-8")
    if args.output_json:
        args.output_json.parent.mkdir(parents=True, exist_ok=True)
        args.output_json.write_text(serialized + "\n", encoding="utf-8")
    print(json.dumps(snapshot["metadata"]["scope"], ensure_ascii=False))


if __name__ == "__main__":
    main()
